from pathlib import Path
from io import BytesIO
from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree as E
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path.cwd();W='http://schemas.openxmlformats.org/wordprocessingml/2006/main';R='http://schemas.openxmlformats.org/officeDocument/2006/relationships';REL='http://schemas.openxmlformats.org/package/2006/relationships';CT='http://schemas.openxmlformats.org/package/2006/content-types'
doc=Document();h=doc.sections[0].first_page_header
p=h.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(10)
r=p.add_run();pic=r.add_picture(str(ROOT/'assets/business_card.jpg'),width=Pt(70),height=Pt(60))
blip=pic._inline.xpath('.//pic:blipFill')[0];crop=OxmlElement('a:srcRect')
for k,v in dict(l=67619,t=16667,r=6190,b=44167).items():crop.set(k,str(v))
blip.insert(1,crop)
for field,value,size,font in [('name','NEST & NOOK PROPERTY CARE',11,'Arial'),('registration','202603156166 (KT0615852-M)',12,'Arial'),('address','#16-03, Trellis Residences, 80100, J.B, Johor.',10.5,'Times New Roman'),('phone','TEL NO.: +60111-3380335',11.5,'Times New Roman')]:
 p=h.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1
 r=p.add_run(value);r.bold=True;r.font.name=font;r.font.size=Pt(size)
 # Stable semantic bookmark used by both Word filling and the PDF map builder.
 start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(len(h.paragraphs)));start.set(qn('w:name'),'NestHeader_'+field);p._p.insert(1,start)
 end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(len(h.paragraphs)));p._p.append(end)
p=h.add_paragraph();p.paragraph_format.space_after=Pt(12);p.paragraph_format.line_spacing=Pt(1);p.add_run().font.size=Pt(1)
borders=OxmlElement('w:pBdr');bottom=OxmlElement('w:bottom')
for k,v in dict(val='single',sz='4',space='1',color='000000').items():bottom.set(qn('w:'+k),v)
borders.append(bottom);p._p.get_or_add_pPr().append(borders)
b=BytesIO();doc.save(b)
with ZipFile(b) as hz:
 header=hz.read('word/header1.xml');rels=hz.read('word/_rels/header1.xml.rels').replace(b'media/image1.jpg',b'media/nest-header-logo.jpg');logo=hz.read('word/media/image1.jpg')
for name in ['1. Move in Form - quantity.docx','2. House Rules.docx','3. Room TA_AC.docx','3. Room TA_NOAC.docx']:
 path=ROOT/'agreements'/name
 with ZipFile(path) as z:parts={i.filename:z.read(i.filename) for i in z.infolist()}
 if 'word/nest-header.xml' in parts:raise SystemExit('Header already exists: '+name)
 root=E.fromstring(parts['word/document.xml']);sec=root.find('{'+W+'}body/{'+W+'}sectPr')
 ref=E.Element('{'+W+'}headerReference');ref.set('{'+W+'}type','first');ref.set('{'+R+'}id','rIdNestHeader');sec.insert(0,ref)
 sec.append(E.Element('{'+W+'}titlePg'));sec.find('{'+W+'}pgMar').set('{'+W+'}header','1300')
 parts['word/document.xml']=E.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
 rel=E.fromstring(parts['word/_rels/document.xml.rels']);E.SubElement(rel,'{'+REL+'}Relationship',Id='rIdNestHeader',Type=R+'/header',Target='nest-header.xml');parts['word/_rels/document.xml.rels']=E.tostring(rel,xml_declaration=True,encoding='UTF-8',standalone=True)
 ct=E.fromstring(parts['[Content_Types].xml']);E.SubElement(ct,'{'+CT+'}Override',PartName='/word/nest-header.xml',ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml')
 if not any(n.get('Extension')=='jpg' for n in ct):E.SubElement(ct,'{'+CT+'}Default',Extension='jpg',ContentType='image/jpeg')
 parts['[Content_Types].xml']=E.tostring(ct,xml_declaration=True,encoding='UTF-8',standalone=True)
 parts.update({'word/nest-header.xml':header,'word/_rels/nest-header.xml.rels':rels,'word/media/nest-header-logo.jpg':logo})
 with ZipFile(path,'w',ZIP_DEFLATED) as z:
  for n,v in parts.items():z.writestr(n,v)
 print(name)
